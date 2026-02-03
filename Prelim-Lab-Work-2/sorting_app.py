"""
Modern Sorting Application
Supports: .txt, .xls files
Algorithms: Bubble Sort, Insertion Sort, Merge Sort
Features: Time complexity display, sorting time, progress tracking, auto-load
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import time
import threading
import os

# Try to import openpyxl and xlrd for Excel support
try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import xlrd
    XLRD_AVAILABLE = True
except ImportError:
    XLRD_AVAILABLE = False


class ProgressDialog(tk.Toplevel):
    """Modern progress dialog that mirrors actual sorting progress"""
    
    def __init__(self, parent, title="Sorting in Progress"):
        super().__init__(parent)
        self.title(title)
        self.geometry("400x180")
        self.resizable(False, False)
        self.configure(bg="#1e1e2e")
        
        # Center the dialog
        self.transient(parent)
        self.grab_set()
        
        # Center on parent
        self.update_idletasks()
        x = parent.winfo_x() + (parent.winfo_width() - 400) // 2
        y = parent.winfo_y() + (parent.winfo_height() - 180) // 2
        self.geometry(f"+{x}+{y}")
        
        # Status label
        self.status_label = tk.Label(
            self, 
            text="Initializing...", 
            font=("Segoe UI", 11),
            bg="#1e1e2e", 
            fg="#cdd6f4"
        )
        self.status_label.pack(pady=(25, 10))
        
        # Progress bar style
        style = ttk.Style()
        style.configure(
            "Custom.Horizontal.TProgressbar",
            troughcolor="#313244",
            background="#89b4fa",
            thickness=25
        )
        
        # Progress bar
        self.progress = ttk.Progressbar(
            self, 
            length=350, 
            mode='determinate',
            style="Custom.Horizontal.TProgressbar"
        )
        self.progress.pack(pady=10)
        
        # Percentage label
        self.percent_label = tk.Label(
            self, 
            text="0%", 
            font=("Segoe UI", 12, "bold"),
            bg="#1e1e2e", 
            fg="#a6e3a1"
        )
        self.percent_label.pack(pady=5)
        
        # Time elapsed label
        self.time_label = tk.Label(
            self, 
            text="Time: 0.00s", 
            font=("Segoe UI", 9),
            bg="#1e1e2e", 
            fg="#6c7086"
        )
        self.time_label.pack(pady=5)
        
        self.protocol("WM_DELETE_WINDOW", lambda: None)  # Disable close button
        
    def update_progress(self, value, status="", elapsed_time=0):
        """Update progress bar and labels"""
        self.progress['value'] = value
        self.percent_label.config(text=f"{int(value)}%")
        if status:
            self.status_label.config(text=status)
        self.time_label.config(text=f"Time: {elapsed_time:.3f}s")
        self.update()


class ComparisonDialog(tk.Toplevel):
    """Modern dialog to display algorithm comparison results"""
    
    def __init__(self, parent, results, data_size):
        super().__init__(parent)
        self.title("📊 Algorithm Comparison Results")
        self.geometry("600x450")
        self.resizable(False, False)
        self.configure(bg="#1e1e2e")
        
        # Center the dialog
        self.transient(parent)
        self.grab_set()
        
        # Center on parent
        self.update_idletasks()
        x = parent.winfo_x() + (parent.winfo_width() - 600) // 2
        y = parent.winfo_y() + (parent.winfo_height() - 450) // 2
        self.geometry(f"+{x}+{y}")
        
        # Title
        title_label = tk.Label(
            self,
            text="📊 Algorithm Comparison Results",
            font=("Segoe UI", 16, "bold"),
            bg="#1e1e2e",
            fg="#b4befe"
        )
        title_label.pack(pady=(20, 5))
        
        # Data size info
        size_label = tk.Label(
            self,
            text=f"Dataset Size: {data_size:,} elements",
            font=("Segoe UI", 11),
            bg="#1e1e2e",
            fg="#a6adc8"
        )
        size_label.pack(pady=(0, 20))
        
        # Results container
        results_frame = tk.Frame(self, bg="#1e1e2e")
        results_frame.pack(fill=tk.BOTH, expand=True, padx=30)
        
        # Sort results by time to determine ranking
        sorted_results = sorted(results.items(), key=lambda x: x[1]['time'])
        
        # Color codes for rankings
        rank_colors = ["#a6e3a1", "#f9e2af", "#fab387"]  # green, yellow, orange
        
        for idx, (algo, data) in enumerate(sorted_results):
            # Card for each algorithm
            card = tk.Frame(results_frame, bg="#313244", highlightthickness=1, highlightbackground="#45475a")
            card.pack(fill=tk.X, pady=8)
            
            # Rank badge
            rank_frame = tk.Frame(card, bg=rank_colors[idx], width=40)
            rank_frame.pack(side=tk.LEFT, fill=tk.Y)
            rank_frame.pack_propagate(False)
            
            rank_label = tk.Label(
                rank_frame,
                text=f"#{idx + 1}",
                font=("Segoe UI", 14, "bold"),
                bg=rank_colors[idx],
                fg="#1e1e2e"
            )
            rank_label.pack(expand=True)
            
            # Algorithm info
            info_frame = tk.Frame(card, bg="#313244")
            info_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=15, pady=12)
            
            algo_label = tk.Label(
                info_frame,
                text=algo,
                font=("Segoe UI", 13, "bold"),
                bg="#313244",
                fg="#89b4fa"
            )
            algo_label.pack(anchor=tk.W)
            
            complexity_label = tk.Label(
                info_frame,
                text=f"Time Complexity: {data['complexity']}",
                font=("Segoe UI", 9),
                bg="#313244",
                fg="#6c7086"
            )
            complexity_label.pack(anchor=tk.W)
            
            # Time display
            time_frame = tk.Frame(card, bg="#313244")
            time_frame.pack(side=tk.RIGHT, padx=20)
            
            time_value = tk.Label(
                time_frame,
                text=f"{data['time']:.6f}s",
                font=("Segoe UI", 16, "bold"),
                bg="#313244",
                fg=rank_colors[idx]
            )
            time_value.pack()
            
            # Speed comparison (relative to slowest)
            slowest_time = sorted_results[-1][1]['time']
            if slowest_time > 0 and data['time'] > 0:
                speed_factor = slowest_time / data['time']
                if speed_factor > 1.01:
                    speed_text = f"{speed_factor:.1f}x faster"
                else:
                    speed_text = "baseline"
            else:
                speed_text = "-"
                
            speed_label = tk.Label(
                time_frame,
                text=speed_text,
                font=("Segoe UI", 9),
                bg="#313244",
                fg="#a6adc8"
            )
            speed_label.pack()
        
        # Winner announcement
        winner = sorted_results[0][0]
        winner_time = sorted_results[0][1]['time']
        
        winner_frame = tk.Frame(self, bg="#1e1e2e")
        winner_frame.pack(pady=20)
        
        winner_label = tk.Label(
            winner_frame,
            text=f"🏆 Winner: {winner} ({winner_time:.6f}s)",
            font=("Segoe UI", 12, "bold"),
            bg="#1e1e2e",
            fg="#a6e3a1"
        )
        winner_label.pack()
        
        # Close button
        close_btn = ttk.Button(
            self,
            text="Close",
            style='Modern.TButton',
            command=self.destroy
        )
        close_btn.pack(pady=(0, 20))


class SortingAlgorithms:
    """Sorting algorithms with progress callback support"""
    
    @staticmethod
    def bubble_sort(arr, progress_callback=None):
        """
        Bubble Sort
        Time Complexity: O(n²) - Best: O(n), Average: O(n²), Worst: O(n²)
        Space Complexity: O(1)
        """
        n = len(arr)
        arr = arr.copy()
        total_ops = n * (n - 1) // 2
        current_op = 0
        
        for i in range(n):
            swapped = False
            for j in range(0, n - i - 1):
                if arr[j] > arr[j + 1]:
                    arr[j], arr[j + 1] = arr[j + 1], arr[j]
                    swapped = True
                current_op += 1
                if progress_callback and current_op % max(1, total_ops // 100) == 0:
                    progress_callback(min(100, (current_op / total_ops) * 100))
            if not swapped:
                break
                
        if progress_callback:
            progress_callback(100)
        return arr
    
    @staticmethod
    def insertion_sort(arr, progress_callback=None):
        """
        Insertion Sort
        Time Complexity: O(n²) - Best: O(n), Average: O(n²), Worst: O(n²)
        Space Complexity: O(1)
        """
        n = len(arr)
        arr = arr.copy()
        total_ops = n * (n - 1) // 2
        current_op = 0
        
        for i in range(1, n):
            key = arr[i]
            j = i - 1
            while j >= 0 and arr[j] > key:
                arr[j + 1] = arr[j]
                j -= 1
                current_op += 1
                if progress_callback and current_op % max(1, total_ops // 100) == 0:
                    progress_callback(min(100, (current_op / total_ops) * 100))
            arr[j + 1] = key
            current_op += 1
            
        if progress_callback:
            progress_callback(100)
        return arr
    
    @staticmethod
    def merge_sort(arr, progress_callback=None):
        """
        Merge Sort (Iterative/Bottom-up)
        Time Complexity: O(n log n) - Best: O(n log n), Average: O(n log n), Worst: O(n log n)
        Space Complexity: O(n)
        """
        arr = list(arr)  # Make a copy
        n = len(arr)
        if n <= 1:
            if progress_callback:
                progress_callback(100)
            return arr
        
        # Bottom-up iterative merge sort to avoid recursion limits
        # Calculate total operations for progress
        import math
        total_passes = math.ceil(math.log2(n)) if n > 1 else 1
        current_pass = 0
        
        # Start with sublists of size 1 and double each pass
        size = 1
        while size < n:
            # Merge sublists
            for start in range(0, n, size * 2):
                mid = min(start + size, n)
                end = min(start + size * 2, n)
                
                # Merge arr[start:mid] and arr[mid:end]
                left = arr[start:mid]
                right = arr[mid:end]
                
                merged = []
                i = j = 0
                while i < len(left) and j < len(right):
                    if left[i] <= right[j]:
                        merged.append(left[i])
                        i += 1
                    else:
                        merged.append(right[j])
                        j += 1
                merged.extend(left[i:])
                merged.extend(right[j:])
                
                # Copy back to arr
                arr[start:start + len(merged)] = merged
            
            size *= 2
            current_pass += 1
            
            if progress_callback:
                progress_callback(min(100, (current_pass / total_passes) * 100))
        
        if progress_callback:
            progress_callback(100)
        return arr


class ModernSortingApp:
    """Modern UI Sorting Application"""
    
    # Time complexity information
    COMPLEXITY_INFO = {
        "Bubble Sort": {
            "best": "O(n)",
            "average": "O(n²)",
            "worst": "O(n²)",
            "space": "O(1)",
            "description": "Simple comparison-based algorithm. Best for small datasets."
        },
        "Insertion Sort": {
            "best": "O(n)",
            "average": "O(n²)",
            "worst": "O(n²)",
            "space": "O(1)",
            "description": "Efficient for small or nearly sorted datasets."
        },
        "Merge Sort": {
            "best": "O(n log n)",
            "average": "O(n log n)",
            "worst": "O(n log n)",
            "space": "O(n)",
            "description": "Divide-and-conquer algorithm. Consistent performance."
        }
    }
    
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("🔢 Modern Sorting Application")
        self.root.geometry("1000x700")
        self.root.minsize(900, 600)
        self.root.configure(bg="#1e1e2e")
        
        # Optimize window resizing performance
        self.root.update_idletasks()
        self._resize_job = None
        
        # Data storage
        self.original_data = []
        self.sorted_data = []
        self.current_file = None
        
        # Apply modern styling
        self.setup_styles()
        self.create_ui()
        
        # Bind optimized resize handler
        self.root.bind('<Configure>', self._on_resize)
        
        # Auto-load last file if exists
        self.auto_load_file()
    
    def _on_resize(self, event):
        """Debounced resize handler to prevent lag during window resize"""
        # Only handle root window resize events
        if event.widget != self.root:
            return
        
        # Cancel previous scheduled update
        if self._resize_job:
            self.root.after_cancel(self._resize_job)
        
        # Schedule update after brief delay (debounce)
        self._resize_job = self.root.after(50, self._handle_resize)
    
    def _handle_resize(self):
        """Handle the actual resize after debounce delay"""
        self._resize_job = None
        # Force update of geometry
        self.root.update_idletasks()
        
    def setup_styles(self):
        """Configure modern ttk styles"""
        style = ttk.Style()
        style.theme_use('clam')
        
        # Configure colors (Catppuccin Mocha theme)
        colors = {
            'bg': '#1e1e2e',
            'surface': '#313244',
            'overlay': '#45475a',
            'text': '#cdd6f4',
            'subtext': '#a6adc8',
            'blue': '#89b4fa',
            'green': '#a6e3a1',
            'red': '#f38ba8',
            'yellow': '#f9e2af',
            'lavender': '#b4befe'
        }
        
        # Frame style
        style.configure('Modern.TFrame', background=colors['bg'])
        style.configure('Card.TFrame', background=colors['surface'])
        
        # Label style
        style.configure('Modern.TLabel', 
                       background=colors['bg'], 
                       foreground=colors['text'],
                       font=('Segoe UI', 10))
        style.configure('Title.TLabel', 
                       background=colors['bg'], 
                       foreground=colors['lavender'],
                       font=('Segoe UI', 16, 'bold'))
        style.configure('Card.TLabel', 
                       background=colors['surface'], 
                       foreground=colors['text'],
                       font=('Segoe UI', 10))
        style.configure('CardTitle.TLabel', 
                       background=colors['surface'], 
                       foreground=colors['blue'],
                       font=('Segoe UI', 11, 'bold'))
        
        # Button style
        style.configure('Modern.TButton',
                       background=colors['blue'],
                       foreground='#1e1e2e',
                       font=('Segoe UI', 10, 'bold'),
                       padding=(15, 8))
        style.map('Modern.TButton',
                 background=[('active', colors['lavender']), ('pressed', colors['overlay'])])
        
        # Combobox style
        style.configure('Modern.TCombobox',
                       background=colors['surface'],
                       foreground=colors['text'],
                       fieldbackground=colors['surface'],
                       font=('Segoe UI', 10))
        
        # Radiobutton style
        style.configure('Modern.TRadiobutton',
                       background=colors['surface'],
                       foreground=colors['text'],
                       font=('Segoe UI', 10))
        
    def create_ui(self):
        """Create the modern user interface"""
        # Main container
        main_frame = ttk.Frame(self.root, style='Modern.TFrame')
        main_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=20)
        
        # Title
        title_label = ttk.Label(
            main_frame, 
            text="🔢 Modern Sorting Application",
            style='Title.TLabel'
        )
        title_label.pack(pady=(0, 15))
        
        # Top section - File Input
        file_card = self.create_card(main_frame, "📁 File Input")
        file_card.pack(fill=tk.X, pady=(0, 10))
        
        file_inner = ttk.Frame(file_card, style='Card.TFrame')
        file_inner.pack(fill=tk.X, padx=15, pady=10)
        
        self.file_label = ttk.Label(
            file_inner, 
            text="No file selected", 
            style='Card.TLabel',
            width=30
        )
        self.file_label.pack(side=tk.LEFT, padx=(0, 10))
        
        browse_btn = ttk.Button(
            file_inner, 
            text="📂 Browse",
            style='Modern.TButton',
            command=self.browse_file
        )
        browse_btn.pack(side=tk.LEFT, padx=5)
        
        reload_btn = ttk.Button(
            file_inner, 
            text="🔄",
            style='Modern.TButton',
            command=self.reload_file,
            width=3
        )
        reload_btn.pack(side=tk.LEFT, padx=5)
        
        # Sorting options card - separate row
        sort_card = self.create_card(main_frame, "⚙️ Sorting Options")
        sort_card.pack(fill=tk.X, pady=(0, 10))
        
        # First row - Algorithm and Dataset
        sort_row1 = ttk.Frame(sort_card, style='Card.TFrame')
        sort_row1.pack(fill=tk.X, padx=15, pady=(10, 5))
        
        # Algorithm dropdown
        algo_label = tk.Label(
            sort_row1,
            text="Algorithm:",
            font=("Segoe UI", 10),
            bg="#313244",
            fg="#cdd6f4"
        )
        algo_label.pack(side=tk.LEFT, padx=(0, 5))
        
        self.sort_var = tk.StringVar(value="Bubble Sort")
        algorithms = ["Bubble Sort", "Insertion Sort", "Merge Sort"]
        
        self.algo_dropdown = ttk.Combobox(
            sort_row1,
            textvariable=self.sort_var,
            values=algorithms,
            state="readonly",
            style='Modern.TCombobox',
            width=14
        )
        self.algo_dropdown.pack(side=tk.LEFT, padx=(0, 15))
        self.algo_dropdown.bind("<<ComboboxSelected>>", lambda e: self.update_complexity_display())
        
        # Dataset size input for random generation
        size_label = tk.Label(
            sort_row1,
            text="Size:",
            font=("Segoe UI", 10),
            bg="#313244",
            fg="#cdd6f4"
        )
        size_label.pack(side=tk.LEFT, padx=(0, 5))
        
        self.size_var = tk.StringVar(value="100")
        self.size_entry = ttk.Entry(
            sort_row1,
            textvariable=self.size_var,
            width=8,
            font=("Segoe UI", 10)
        )
        self.size_entry.pack(side=tk.LEFT, padx=(0, 5))
        
        generate_btn = ttk.Button(
            sort_row1,
            text="🎲 Generate",
            style='Modern.TButton',
            command=self.generate_random_data
        )
        generate_btn.pack(side=tk.LEFT)
        
        # Second row - Action buttons
        sort_row2 = ttk.Frame(sort_card, style='Card.TFrame')
        sort_row2.pack(fill=tk.X, padx=15, pady=(0, 10))
        
        # Run button
        run_btn = ttk.Button(
            sort_row2, 
            text="▶ Run",
            style='Modern.TButton',
            command=self.start_sorting
        )
        run_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        compare_btn = ttk.Button(
            sort_row2, 
            text="📊 Compare All",
            style='Modern.TButton',
            command=self.compare_algorithms
        )
        compare_btn.pack(side=tk.LEFT)
        
        # Complexity info card
        complexity_card = self.create_card(main_frame, "📊 Time Complexity Analysis")
        complexity_card.pack(fill=tk.X, pady=(0, 15))
        
        complexity_inner = ttk.Frame(complexity_card, style='Card.TFrame')
        complexity_inner.pack(fill=tk.X, padx=15, pady=10)
        
        # Complexity labels
        self.best_label = self.create_complexity_item(complexity_inner, "Best Case", "O(n)")
        self.avg_label = self.create_complexity_item(complexity_inner, "Average Case", "O(n²)")
        self.worst_label = self.create_complexity_item(complexity_inner, "Worst Case", "O(n²)")
        self.space_label = self.create_complexity_item(complexity_inner, "Space", "O(1)")
        
        self.desc_label = ttk.Label(
            complexity_card,
            text="Simple comparison-based algorithm. Best for small datasets.",
            style='Card.TLabel',
            wraplength=800
        )
        self.desc_label.pack(padx=15, pady=(0, 10))
        
        # Data display section
        data_frame = ttk.Frame(main_frame, style='Modern.TFrame')
        data_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 15))
        
        # Original data card
        orig_card = self.create_card(data_frame, "📥 Original Data")
        orig_card.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 10))
        
        self.original_text = self.create_text_widget(orig_card)
        self.original_text.pack(fill=tk.BOTH, expand=True, padx=15, pady=(0, 15))
        
        # Sorted data card
        sorted_card = self.create_card(data_frame, "📤 Sorted Data")
        sorted_card.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        self.sorted_text = self.create_text_widget(sorted_card)
        self.sorted_text.pack(fill=tk.BOTH, expand=True, padx=15, pady=(0, 15))
        
        # Stats card
        stats_card = self.create_card(main_frame, "📈 Statistics")
        stats_card.pack(fill=tk.X)
        
        stats_inner = ttk.Frame(stats_card, style='Card.TFrame')
        stats_inner.pack(fill=tk.X, padx=15, pady=10)
        
        self.stats_label = tk.Label(
            stats_inner,
            text="Load a file to see statistics",
            font=("Segoe UI", 11),
            bg="#313244",
            fg="#a6adc8"
        )
        self.stats_label.pack(side=tk.LEFT)
        
        self.time_label = tk.Label(
            stats_inner,
            text="",
            font=("Segoe UI", 11, "bold"),
            bg="#313244",
            fg="#a6e3a1"
        )
        self.time_label.pack(side=tk.RIGHT)
        
        # Save button
        save_btn = ttk.Button(
            stats_inner, 
            text="💾 Save Sorted",
            style='Modern.TButton',
            command=self.save_sorted_data
        )
        save_btn.pack(side=tk.RIGHT, padx=20)
        
    def create_card(self, parent, title):
        """Create a modern card widget"""
        card = tk.Frame(parent, bg="#313244", highlightthickness=1, highlightbackground="#45475a")
        
        title_label = tk.Label(
            card,
            text=title,
            font=("Segoe UI", 11, "bold"),
            bg="#313244",
            fg="#89b4fa"
        )
        title_label.pack(anchor=tk.W, padx=15, pady=(10, 5))
        
        return card
    
    def create_complexity_item(self, parent, label, value):
        """Create a complexity display item"""
        frame = tk.Frame(parent, bg="#313244")
        frame.pack(side=tk.LEFT, padx=20)
        
        lbl = tk.Label(
            frame,
            text=label,
            font=("Segoe UI", 9),
            bg="#313244",
            fg="#6c7086"
        )
        lbl.pack()
        
        val = tk.Label(
            frame,
            text=value,
            font=("Segoe UI", 14, "bold"),
            bg="#313244",
            fg="#f9e2af"
        )
        val.pack()
        
        return val
    
    def create_text_widget(self, parent):
        """Create a styled text widget with scrollbar"""
        frame = tk.Frame(parent, bg="#313244")
        
        scrollbar = ttk.Scrollbar(frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        text = tk.Text(
            frame,
            bg="#1e1e2e",
            fg="#cdd6f4",
            font=("Consolas", 10),
            insertbackground="#cdd6f4",
            selectbackground="#45475a",
            wrap=tk.NONE,  # Use NONE for better resize performance
            yscrollcommand=scrollbar.set,
            relief=tk.FLAT,
            padx=10,
            pady=10,
            undo=False,  # Disable undo for performance
            maxundo=0
        )
        text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.config(command=text.yview)
        
        # Add horizontal scrollbar for long content
        h_scrollbar = ttk.Scrollbar(frame, orient=tk.HORIZONTAL, command=text.xview)
        h_scrollbar.pack(side=tk.BOTTOM, fill=tk.X)
        text.config(xscrollcommand=h_scrollbar.set)
        
        return frame
    
    def update_complexity_display(self):
        """Update complexity information based on selected algorithm"""
        algo = self.sort_var.get()
        info = self.COMPLEXITY_INFO[algo]
        
        self.best_label.config(text=info["best"])
        self.avg_label.config(text=info["average"])
        self.worst_label.config(text=info["worst"])
        self.space_label.config(text=info["space"])
        self.desc_label.config(text=info["description"])
    
    def browse_file(self):
        """Open file browser dialog"""
        filetypes = [
            ("Supported files", "*.txt *.xls *.xlsx"),
            ("Text files", "*.txt"),
            ("Excel files", "*.xls *.xlsx"),
            ("All files", "*.*")
        ]
        
        filepath = filedialog.askopenfilename(
            title="Select a file",
            filetypes=filetypes
        )
        
        if filepath:
            self.load_file(filepath)
    
    def load_file(self, filepath):
        """Load data from file"""
        try:
            ext = os.path.splitext(filepath)[1].lower()
            data = []
            
            if ext == '.txt':
                data = self.load_txt_file(filepath)
            elif ext in ['.xls', '.xlsx']:
                data = self.load_excel_file(filepath, ext)
            else:
                messagebox.showerror("Error", f"Unsupported file format: {ext}")
                return
            
            if not data:
                messagebox.showwarning("Warning", "No numeric data found in file!")
                return
            
            self.original_data = data
            self.current_file = filepath
            self.sorted_data = []
            
            # Update UI
            filename = os.path.basename(filepath)
            self.file_label.config(text=f"📄 {filename}")
            
            # Display original data
            self.display_data(self.original_text, data)
            
            # Clear sorted data
            sorted_text = self.sorted_text.winfo_children()[1]  # Get text widget
            sorted_text.delete(1.0, tk.END)
            sorted_text.insert(tk.END, "Click 'Sort Data' to sort...")
            
            # Update stats
            self.stats_label.config(
                text=f"📊 Elements: {len(data)} | Min: {min(data)} | Max: {max(data)} | File: {filename}"
            )
            self.time_label.config(text="")
            
            # Save filepath for auto-load
            self.save_last_file(filepath)
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load file: {str(e)}")
    
    def load_txt_file(self, filepath):
        """Load numbers from text file"""
        data = []
        with open(filepath, 'r') as f:
            content = f.read()
            # Try to parse numbers separated by various delimiters
            import re
            numbers = re.findall(r'-?\d+\.?\d*', content)
            for num in numbers:
                try:
                    if '.' in num:
                        data.append(float(num))
                    else:
                        data.append(int(num))
                except ValueError:
                    continue
        return data
    
    def load_excel_file(self, filepath, ext):
        """Load numbers from Excel file"""
        data = []
        
        if ext == '.xlsx':
            if not OPENPYXL_AVAILABLE:
                messagebox.showerror(
                    "Missing Dependency",
                    "Please install openpyxl to read .xlsx files:\npip install openpyxl"
                )
                return []
            
            wb = openpyxl.load_workbook(filepath)
            sheet = wb.active
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        try:
                            num = float(cell.value)
                            if num == int(num):
                                data.append(int(num))
                            else:
                                data.append(num)
                        except (ValueError, TypeError):
                            continue
                            
        elif ext == '.xls':
            if not XLRD_AVAILABLE:
                messagebox.showerror(
                    "Missing Dependency",
                    "Please install xlrd to read .xls files:\npip install xlrd"
                )
                return []
            
            wb = xlrd.open_workbook(filepath)
            sheet = wb.sheet_by_index(0)
            for row in range(sheet.nrows):
                for col in range(sheet.ncols):
                    cell_value = sheet.cell_value(row, col)
                    if cell_value:
                        try:
                            num = float(cell_value)
                            if num == int(num):
                                data.append(int(num))
                            else:
                                data.append(num)
                        except (ValueError, TypeError):
                            continue
        
        return data
    
    def display_data(self, text_frame, data):
        """Display data in text widget"""
        text_widget = text_frame.winfo_children()[1]  # Get text widget (after scrollbar)
        text_widget.delete(1.0, tk.END)
        
        # Format data nicely
        formatted = ", ".join(str(x) for x in data)
        text_widget.insert(tk.END, formatted)
    
    def reload_file(self):
        """Reload the current file"""
        if self.current_file and os.path.exists(self.current_file):
            self.load_file(self.current_file)
        else:
            messagebox.showinfo("Info", "No file to reload. Please select a file first.")
    
    def generate_random_data(self):
        """Generate random dataset with user-specified size"""
        import random
        
        try:
            size = int(self.size_var.get())
            if size <= 0:
                messagebox.showwarning("Invalid Size", "Please enter a positive integer.")
                return
            if size > 100000:
                if not messagebox.askyesno("Large Dataset", 
                    f"Generating {size:,} elements may take time.\nContinue?"):
                    return
        except ValueError:
            messagebox.showerror("Invalid Input", "Please enter a valid integer for dataset size.")
            return
        
        # Generate random integers
        data = [random.randint(1, 10000) for _ in range(size)]
        
        self.original_data = data
        self.sorted_data = []
        self.current_file = None
        
        # Update UI
        self.file_label.config(text=f"🎲 Random Data ({size:,} elements)")
        self.display_data(self.original_text, data)
        
        # Clear sorted data
        sorted_text = self.sorted_text.winfo_children()[1]
        sorted_text.delete(1.0, tk.END)
        sorted_text.insert(tk.END, "Click 'Run' to sort...")
        
        # Update stats
        self.stats_label.config(
            text=f"📊 Elements: {len(data):,} | Min: {min(data)} | Max: {max(data)} | Source: Random"
        )
        self.time_label.config(text="")
    
    def verify_sorted(self, data):
        """Verify that the array is correctly sorted in ascending order"""
        for i in range(len(data) - 1):
            if data[i] > data[i + 1]:
                return False
        return True
    
    def start_sorting(self):
        """Start the sorting process in a separate thread"""
        if not self.original_data:
            messagebox.showwarning("Warning", "Please load a file first!")
            return
        
        # Create progress dialog
        self.progress_dialog = ProgressDialog(self.root, "Sorting in Progress")
        
        # Start sorting in a separate thread
        thread = threading.Thread(target=self.sort_data)
        thread.start()
    
    def sort_data(self):
        """Perform the sorting operation"""
        algo = self.sort_var.get()
        data = self.original_data.copy()
        
        start_time = time.time()
        
        def progress_callback(progress):
            elapsed = time.time() - start_time
            self.root.after(0, lambda: self.progress_dialog.update_progress(
                progress, 
                f"Sorting with {algo}...",
                elapsed
            ))
        
        # Select sorting algorithm
        if algo == "Bubble Sort":
            sorted_data = SortingAlgorithms.bubble_sort(data, progress_callback)
        elif algo == "Insertion Sort":
            sorted_data = SortingAlgorithms.insertion_sort(data, progress_callback)
        else:  # Merge Sort
            sorted_data = SortingAlgorithms.merge_sort(data, progress_callback)
        
        end_time = time.time()
        elapsed_time = end_time - start_time
        
        self.sorted_data = sorted_data
        
        # Update UI on main thread
        self.root.after(0, lambda: self.sorting_complete(elapsed_time))
    
    def sorting_complete(self, elapsed_time):
        """Handle sorting completion"""
        # Close progress dialog
        self.progress_dialog.destroy()
        
        # Verify sorting is correct
        is_verified = self.verify_sorted(self.sorted_data)
        verification_status = "✅ Verified" if is_verified else "❌ Verification Failed"
        
        # Display sorted data
        self.display_data(self.sorted_text, self.sorted_data)
        
        # Update time label
        algo = self.sort_var.get()
        complexity = self.COMPLEXITY_INFO[algo]["average"]
        self.time_label.config(
            text=f"⏱️ Sorted in: {elapsed_time:.6f}s | Complexity: {complexity} | {verification_status}"
        )
        
        messagebox.showinfo(
            "Sorting Complete",
            f"✅ Successfully sorted {len(self.sorted_data):,} elements!\n\n"
            f"Algorithm: {algo}\n"
            f"Time: {elapsed_time:.6f} seconds\n"
            f"Time Complexity: {complexity}\n"
            f"Verification: {verification_status}"
        )
    
    def compare_algorithms(self):
        """Compare all sorting algorithms on the current dataset"""
        if not self.original_data:
            messagebox.showwarning("Warning", "Please load a file first!")
            return
        
        # Show a simple waiting message
        self.root.config(cursor="wait")
        self.root.update()
        
        # Run comparison directly (blocking but simple)
        results = self._run_comparison_sync()
        
        self.root.config(cursor="")
        
        # Show comparison dialog
        ComparisonDialog(self.root, results, len(self.original_data))
    
    def _run_comparison_sync(self):
        """Run all sorting algorithms synchronously and collect timing data"""
        data = self.original_data.copy()
        results = {}
        
        algorithms = [
            ("Bubble Sort", SortingAlgorithms.bubble_sort, "O(n²)"),
            ("Insertion Sort", SortingAlgorithms.insertion_sort, "O(n²)"),
            ("Merge Sort", SortingAlgorithms.merge_sort, "O(n log n)")
        ]
        
        for name, sort_func, complexity in algorithms:
            # Time the algorithm (without progress callback for accurate timing)
            start_time = time.time()
            sort_func(data.copy(), None)  # No callback for accurate timing
            end_time = time.time()
            
            elapsed = end_time - start_time
            results[name] = {
                'time': elapsed,
                'complexity': complexity
            }
        
        return results
    
    def save_sorted_data(self):
        """Save sorted data to file"""
        if not self.sorted_data:
            messagebox.showwarning("Warning", "No sorted data to save!")
            return
        
        filepath = filedialog.asksaveasfilename(
            title="Save Sorted Data",
            defaultextension=".txt",
            filetypes=[("Text files", "*.txt"), ("All files", "*.*")]
        )
        
        if filepath:
            try:
                with open(filepath, 'w') as f:
                    f.write(", ".join(str(x) for x in self.sorted_data))
                messagebox.showinfo("Success", f"Sorted data saved to:\n{filepath}")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to save file: {str(e)}")
    
    def save_last_file(self, filepath):
        """Save last file path for auto-load"""
        try:
            config_path = os.path.join(os.path.dirname(__file__), '.sorting_app_config')
            with open(config_path, 'w') as f:
                f.write(filepath)
        except:
            pass
    
    def auto_load_file(self):
        """Auto-load the last used file"""
        try:
            config_path = os.path.join(os.path.dirname(__file__), '.sorting_app_config')
            if os.path.exists(config_path):
                with open(config_path, 'r') as f:
                    filepath = f.read().strip()
                if os.path.exists(filepath):
                    self.load_file(filepath)
        except:
            pass
    
    def run(self):
        """Run the application"""
        self.root.mainloop()


if __name__ == "__main__":
    app = ModernSortingApp()
    app.run()
